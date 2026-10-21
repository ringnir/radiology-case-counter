# 研修手帳テンプレート（xlsx）から、アプリに埋め込む定義を生成する
import openpyxl, base64, json, re, warnings, sys, os
warnings.filterwarnings('ignore')
# 使い方: python3 tools/gen-template.py <様式.xlsx> [--embed]
#   既定: src/template.js（構造のみ）を出力。--embed を付けると src/template-b64.js（様式本体の埋め込み用）も出力
SRC=sys.argv[1]
EMBED='--embed' in sys.argv
OUT=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','src')
wb=openpyxl.load_workbook(SRC)
ws=wb.worksheets[0]
areas=[]; slots=[]
area=None; grp=None; size=1; gidx=0
for r in range(5,ws.max_row+1):
    c=ws.cell(r,3).value
    if not isinstance(c,int): continue
    a=ws.cell(r,1).value; b=ws.cell(r,2).value
    if a:
        area=re.sub(r'\s+','',str(a))
        if area not in areas: areas.append(area)
    if b:
        t=str(b)
        m=re.search(r'（(\d+)症例）',t)
        size=int(m.group(1)) if m else 1
        grp=re.sub(r'\s+','',re.sub(r'（\d+症例）','',t))
        gidx=0
    else:
        gidx+=1
    slots.append([c,r,areas.index(area),grp,gidx,size])
assert [s[0] for s in slots]==list(range(1,101)), 'slot numbering'
print(len(slots),areas)
from collections import Counter
print(Counter(s[2] for s in slots))
rt_sites={}
wr=wb.worksheets[8]
for r in range(6,66):
    h=wr.cell(r,8).value
    if h: rt_sites[r-5]=h
print('rt prefilled',rt_sites)
b64=base64.b64encode(open(SRC,'rb').read()).decode() if EMBED else ''
sheets={'a':{'f':'sheet1.xml','name':wb.worksheets[0].title},
        'xray':{'f':'sheet2.xml','name':wb.worksheets[1].title,'cap':800},
        'gi':{'f':'sheet3.xml','name':wb.worksheets[2].title,'cap':120},
        'us':{'f':'sheet4.xml','name':wb.worksheets[3].title,'cap':240},
        'ct':{'f':'sheet5.xml','name':wb.worksheets[4].title,'cap':1200},
        'mri':{'f':'sheet6.xml','name':wb.worksheets[5].title,'cap':600},
        'rn':{'f':'sheet7.xml','name':wb.worksheets[6].title,'cap':100},
        'ivr':{'f':'sheet8.xml','name':wb.worksheets[7].title,'cap':60},
        'rt':{'f':'sheet9.xml','name':wb.worksheets[8].title,'cap':60}}
out='/* 自動生成（tools/gen-template.py）: 研修手帳様式（形成評価 Ⅳ.経験）の構造（領域・疾患群・シート名・行数）。様式ファイルそのものは含みません */\n'
out+='const TPL_AREAS = '+json.dumps(areas,ensure_ascii=False)+';\n'
out+='const TPL_SLOTS = '+json.dumps(slots,ensure_ascii=False,separators=(',',':'))+';  // [番号, 行, 領域index, 疾患群, 群内index, 群の症例数]\n'
out+='const TPL_SHEETS = '+json.dumps(sheets,ensure_ascii=False)+';\n'
out+='const TPL_RT_SITES = '+json.dumps(rt_sites,ensure_ascii=False)+';  // 放射線治療シートで部位が先に入っている番号\n'
open(os.path.join(OUT,'template.js'),'w',encoding='utf8').write(out)
print('written',len(out))
if EMBED:
    open(os.path.join(OUT,'template-b64.js'),'w',encoding='utf8').write('/* 自動生成: 様式ファイル本体（埋め込み版のみ。公開リポジトリには含めません） */\nconst TPL_B64 = "'+b64+'";\n')
    print('embedded b64 written')
